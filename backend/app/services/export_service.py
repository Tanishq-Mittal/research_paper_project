import io
import json
from typing import Dict, Any, List
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

class ExportService:
    @staticmethod
    def export_to_markdown(title: str, content: str, metadata: Dict[str, Any] = None) -> str:
        md = f"# {title}\n\n"
        if metadata:
            md += "### Document Metadata\n"
            for k, v in metadata.items():
                md += f"- **{k}**: {v}\n"
            md += "\n---\n\n"
        md += content
        return md

    @staticmethod
    def export_to_docx(title: str, content: str, metadata: Dict[str, Any] = None) -> io.BytesIO:
        doc = Document()
        doc.add_heading(title, level=0)
        
        if metadata:
            doc.add_heading("Metadata", level=2)
            for k, v in metadata.items():
                p = doc.add_paragraph()
                p.add_run(f"{k}: ").bold = True
                p.add_run(str(v))
            doc.add_paragraph("--------------------------------------------------")

        lines = content.split("\n")
        for line in lines:
            line_str = line.strip()
            if line_str.startswith("# "):
                doc.add_heading(line_str[2:], level=1)
            elif line_str.startswith("## "):
                doc.add_heading(line_str[3:], level=2)
            elif line_str.startswith("### "):
                doc.add_heading(line_str[4:], level=3)
            elif line_str.startswith("- "):
                doc.add_paragraph(line_str[2:], style='List Bullet')
            elif line_str:
                doc.add_paragraph(line_str)

        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        return file_stream

    @staticmethod
    def export_to_pdf(title: str, content: str, metadata: Dict[str, Any] = None) -> io.BytesIO:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor('#4338ca'),
            spaceAfter=12
        )
        h2_style = ParagraphStyle(
            'CustomH2',
            parent=styles['Heading2'],
            fontSize=14,
            leading=18,
            textColor=colors.HexColor('#1e293b'),
            spaceBefore=10,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'CustomBody',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#334155'),
            spaceAfter=6
        )

        elements = []
        elements.append(Paragraph(title, title_style))
        elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#cbd5e1'), spaceAfter=12))

        if metadata:
            elements.append(Paragraph("<b>Document Information:</b>", body_style))
            for k, v in metadata.items():
                elements.append(Paragraph(f"• <b>{k}:</b> {v}", body_style))
            elements.append(Spacer(1, 10))

        # Parse text lines into flowables
        for line in content.split("\n"):
            line_str = line.strip()
            if not line_str:
                elements.append(Spacer(1, 4))
            elif line_str.startswith("# ") or line_str.startswith("## "):
                clean_h = line_str.lstrip("#").strip()
                elements.append(Paragraph(f"<b>{clean_h}</b>", h2_style))
            elif line_str.startswith("- ") or line_str.startswith("* "):
                clean_bullet = line_str[2:].replace("<", "&lt;").replace(">", "&gt;")
                elements.append(Paragraph(f"&bull; {clean_bullet}", body_style))
            else:
                clean_text = line_str.replace("<", "&lt;").replace(">", "&gt;")
                elements.append(Paragraph(clean_text, body_style))

        doc.build(elements)
        buffer.seek(0)
        return buffer

export_service = ExportService()
